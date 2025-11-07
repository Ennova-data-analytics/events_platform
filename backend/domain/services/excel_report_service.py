from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def create_registrations_excel(registrations_data: list[tuple], event_name: str, status_filter: list[str] = None) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"

    headers = ["Registration ID", "Full Name", "Email", "Degree", "Study Year", "Status"]

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill 
        cell.font = header_font
        cell.alignment = header_alignment
    
    for row_num, reg_data in enumerate(registrations_data, 2):
        (registration_id, full_name, email, degree, study_year, status) = reg_data

        ws.cell(row=row_num, column=1, value=registration_id)
        ws.cell(row=row_num, column=2, value=full_name)
        ws.cell(row=row_num, column=3, value=email)
        ws.cell(row=row_num, column=4, value=degree)
        ws.cell(row=row_num, column=5, value=study_year)
        ws.cell(row=row_num, column=6, value=status)
    
    for col_num in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_num)
        max_length = 0
        for cell in ws[column_letter]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception as _:
                pass 
        
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    metadata_ws = wb.create_sheet("Export Info")
    metadata_ws.cell(row=1, column=1, value="Export Date:").font = Font(bold=True)
    metadata_ws.cell(row=1, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    metadata_ws.cell(row=2, column=1, value="Event:").font = Font(bold=True)
    metadata_ws.cell(row=2, column=2, value=event_name)
        
    metadata_ws.cell(row=3, column=1, value="Total Records:").font = Font(bold=True)
    metadata_ws.cell(row=3, column=2, value=len(registrations_data))
        
    if status_filter:
        metadata_ws.cell(row=4, column=1, value="Status Filter:").font = Font(bold=True)
        metadata_ws.cell(row=4, column=2, value=", ".join(status_filter))
        
    metadata_ws.column_dimensions['A'].width = 15
    metadata_ws.column_dimensions['B'].width = 40

    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return excel_file
    



